"""Disposable conversion process; the API enforces timeout and output limits."""

import sys
import re
from pathlib import Path
from io import BytesIO


def convert(source, target, output):
    import fitz

    with fitz.open(source) as pdf:
        if pdf.needs_pass or not 0 < len(pdf) <= 200:
            raise ValueError("Use an unlocked PDF with 1–200 pages.")
        if target in {"docx", "xlsx", "md"}:
            texts = [page.get_text(sort=True) for page in pdf]
            if not any(text.strip() for text in texts):
                raise ValueError("Run OCR before converting scanned PDFs.")
        if target == "md":
            # Escape punctuation so PDF text cannot turn into active HTML or images.
            def escape(text):
                return re.sub(r"([\\`*_{}\[\]()#+.!<>|~\-])", r"\\\1", text)

            sections = []
            for index, text in enumerate(texts):
                lines = [escape(line) for line in text.splitlines()]
                sections.append(f"## Page {index + 1}\n\n" + "  \n".join(lines))
            Path(output).write_text(
                "\n\n---\n\n".join(sections) + "\n", encoding="utf-8"
            )
        elif target == "docx":
            from docx import Document

            doc = Document()
            for index, text in enumerate(texts):
                if index:
                    doc.add_page_break()
                for line in text.splitlines():
                    doc.add_paragraph(line)
            doc.save(output)
        elif target == "xlsx":
            from openpyxl import Workbook

            book = Workbook()
            book.remove(book.active)
            for index, text in enumerate(texts):
                sheet = book.create_sheet(f"Page {index + 1}")
                sheet.column_dimensions["A"].width = 100
                for row, line in enumerate(text.splitlines(), 1):
                    cell = sheet.cell(row, 1, line)
                    cell.data_type = "s"
            book.save(output)
        elif target == "pptx":
            from pptx import Presentation
            from pptx.util import Inches

            deck = Presentation()
            deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
            for page in pdf:
                factor = min(2, 1600 / max(page.rect.width, page.rect.height))
                pix = page.get_pixmap(matrix=fitz.Matrix(factor, factor))
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                scale = min(
                    deck.slide_width / pix.width, deck.slide_height / pix.height
                )
                w, h = int(pix.width * scale), int(pix.height * scale)
                slide.shapes.add_picture(
                    BytesIO(pix.tobytes("png")),
                    (deck.slide_width - w) // 2,
                    (deck.slide_height - h) // 2,
                    w,
                    h,
                )
            deck.save(output)


if __name__ == "__main__":
    convert(*sys.argv[1:])
