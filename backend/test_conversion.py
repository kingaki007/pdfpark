from pathlib import Path
import tempfile
import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from convert_document import convert


def test_pdf_office_outputs():
    with tempfile.TemporaryDirectory() as folder:
        root = Path(folder)
        source = root / "source.pdf"
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), "=SUM(A1:A2)")
        pdf.new_page().insert_text((72, 72), "Second page")
        pdf.save(source)
        pdf.close()
        for target in ("docx", "xlsx", "pptx"):
            convert(str(source), target, str(root / ("out." + target)))
        assert "Second page" in "\n".join(
            p.text for p in Document(root / "out.docx").paragraphs
        )
        book = load_workbook(root / "out.xlsx")
        assert book.sheetnames == ["Page 1", "Page 2"]
        assert book.worksheets[0]["A1"].data_type == "s"
        assert book.worksheets[0]["A1"].value == "=SUM(A1:A2)"
        assert len(Presentation(root / "out.pptx").slides) == 2


def test_blank_pdf_requires_ocr():
    import pytest

    with tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / "blank.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(source)
        pdf.close()
        for target in ("docx", "xlsx", "md"):
            with pytest.raises(ValueError, match="OCR"):
                convert(str(source), target, str(Path(folder) / ("out." + target)))


def test_conversion_api_rejects_invalid_formats_and_empty_files():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from conversion import router

    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        assert (
            client.post("/api/convert?source=exe&target=pdf", content=b"x").status_code
            == 422
        )
        assert (
            client.post("/api/convert?source=pdf&target=docx", content=b"").status_code
            == 422
        )


def test_office_pdf_route():
    import io
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    import conversion

    document = Document()
    document.add_paragraph("PDF Park ARM conversion check")
    source = io.BytesIO()
    document.save(source)
    app = FastAPI()
    app.include_router(conversion.router)
    with TestClient(app) as client:
        response = client.post(
            "/api/convert?source=docx&target=pdf", content=source.getvalue()
        )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    with fitz.open(stream=response.content, filetype="pdf") as result:
        assert "PDF Park ARM conversion check" in "".join(
            page.get_text() for page in result
        )


def test_markdown_export():
    with tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / "source.pdf"
        out = Path(folder) / "out.md"
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Hello *world* <script>")
            pdf.new_page().insert_text((72, 72), "Second page")
            pdf.save(source)
        convert(str(source), "md", str(out))
        text = out.read_text(encoding="utf-8")
        assert "## Page 1" in text and "## Page 2" in text
        assert r"Hello \*world\* \<script\>" in text
        assert "Second page" in text


def test_markdown_api():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from conversion import router

    app = FastAPI()
    app.include_router(router)
    with fitz.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Markdown export check")
        data = pdf.tobytes()
    with TestClient(app) as client:
        response = client.post("/api/convert?source=pdf&target=md", content=data)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/markdown")
        assert "Markdown export check" in response.text
